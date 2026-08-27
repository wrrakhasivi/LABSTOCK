import openpyxl, json
from openpyxl.utils import get_column_letter

wbv = openpyxl.load_workbook("reagen.xlsx", data_only=True)

def extract_month(sheet, saldo_label):
    ws = wbv[sheet]
    rows=[]
    for r in range(4, 200):
        name = ws.cell(row=r, column=3).value  # C
        if not name: continue
        def g(col): return ws.cell(row=r, column=col).value
        rows.append({
            "reagen": str(name).strip(),
            "qty_per_kit": g(4),
            "avg_2022": g(5),
            "avg_2023": g(6),
            "saldo_awal": g(7),   # G
            "qc": g(39),          # AM
            "total_pemakaian": g(40),  # AN
            "sisa_stock": g(41),  # AO
            "buffer_stock": g(42),# AP
            "satuan": g(43),      # AQ
        })
    return rows

data = {
  "Juli2026": extract_month("Juli2026","Juni2026"),
  "Aug2026": extract_month("Aug2026","Juli2026"),
}
# daily usage from PQ sheets
def extract_pq(sheet):
    ws=wbv[sheet]
    out=[]
    for r in range(2, ws.max_row+1):
        name=ws.cell(row=r,column=1).value
        if not name: continue
        days={}
        for c in range(2,33):  # B..AF = day1..31
            v=ws.cell(row=r,column=c).value
            if v: days[str(c-1)]=v
        if days:
            out.append({"test":str(name).strip(),"days":days})
    return out
data["PQ_Juli"]=extract_pq("PQ_Juli")
data["PQ_Aug"]=extract_pq("PQ_Aug")

# mapping
ws=wbv['Mapping_Test']
mapping=[]
for r in range(2,ws.max_row+1):
    a=ws.cell(row=r,column=1).value
    if not a: continue
    mapping.append({"lis_name":str(a).strip(),
                    "reagen":(str(ws.cell(row=r,column=2).value).strip() if ws.cell(row=r,column=2).value else None),
                    "status":ws.cell(row=r,column=3).value})
data["Mapping_Test"]=mapping

# LIS raw (date-based): period YYYY-MM inferred from sheet; grup carried down
def extract_lis(sheet, default_period):
    ws=wbv[sheet]
    out=[]
    cur_grup=None
    for r in range(2, ws.max_row+1):
        grup=ws.cell(row=r,column=1).value
        name=ws.cell(row=r,column=2).value
        if grup and str(grup).strip():
            cur_grup=str(grup).strip()
        if not name: continue
        src=ws.cell(row=r,column=35).value  # AI SourceFile
        total=ws.cell(row=r,column=34).value  # AH Total
        days={}
        for c in range(3,34):  # C..AG = day1..31
            v=ws.cell(row=r,column=c).value
            if v: days[str(c-2)]=v
        out.append({"period":default_period,"grup":cur_grup,"nama_test":str(name).strip(),
                    "total":total,"source_file":src,"days":days})
    return out
data["LIS_Juli"]=extract_lis("LIS_Juli","2026-07")
try:
    data["LIS_Aug"]=extract_lis("LIS_Aug","2026-08")
except Exception:
    data["LIS_Aug"]=[]

# PRF & Penerimaan from monthly sheets (AR/AS/AT/AU = PR; AV/AW = terima; D=qty/kit)
def extract_prf_penerimaan(sheet, period):
    ws=wbv[sheet]
    prf=[]; pen=[]
    for r in range(4,200):
        name=ws.cell(row=r,column=3).value
        if not name: continue
        name=str(name).strip()
        qpk=ws.cell(row=r,column=4).value
        tgl_pr1=ws.cell(row=r,column=44).value; kit1=ws.cell(row=r,column=45).value
        tgl_pr2=ws.cell(row=r,column=46).value; kit2=ws.cell(row=r,column=47).value
        terima1=ws.cell(row=r,column=48).value
        terima2=ws.cell(row=r,column=49).value
        def isdate(x):
            import datetime as _dt
            return isinstance(x,(_dt.datetime,_dt.date))
        if isdate(tgl_pr1):
            prf.append({"period":period,"reagen":name,"tanggal_pr":tgl_pr1,"reagent_no":1,"kits":kit1 or 1})
        if isdate(tgl_pr2):
            prf.append({"period":period,"reagen":name,"tanggal_pr":tgl_pr2,"reagent_no":2,"kits":kit2 or 1})
        if isdate(terima1):
            pen.append({"period":period,"reagen":name,"tanggal_terima":terima1,"reagent_no":1,"kits":kit1 or 1,"qty_per_kit":qpk})
        if isdate(terima2):
            pen.append({"period":period,"reagen":name,"tanggal_terima":terima2,"reagent_no":2,"kits":kit2 or 1,"qty_per_kit":qpk})
    return prf,pen

pj,nj=extract_prf_penerimaan("Juli2026","2026-07")
pa,na=extract_prf_penerimaan("Aug2026","2026-08")
data["PRF"]=pj+pa
data["Penerimaan"]=nj+na

json.dump(data, open("seed_data.json","w"), default=str, indent=1)
print("LIS_Juli rows:", len(data["LIS_Juli"]))
print("LIS_Aug rows:", len(data["LIS_Aug"]))
print("PRF rows:", len(data["PRF"]))
print("Penerimaan rows:", len(data["Penerimaan"]))
print("Juli reagen:", len(data["Juli2026"]))
print("Aug reagen:", len(data["Aug2026"]))
print("PQ_Juli tests:", len(data["PQ_Juli"]))
print("Mapping:", len(mapping))
print("\nSample Juli rows:")
for x in data["Juli2026"][:5]: print(x)
print("\nSample mapping OK count:", sum(1 for m in mapping if m['status']=='OK'))
print("Sample PQ_Juli:", data["PQ_Juli"][:2])
