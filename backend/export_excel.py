"""Ekspor Pemantauan Stok bulanan ke Excel (kolom harian 1-31 + warna status)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

STATUS_FILL = {
    'critical': PatternFill('solid', fgColor='F8B4B4'),
    'warning': PatternFill('solid', fgColor='FDE68A'),
    'safe': PatternFill('solid', fgColor='BBF7D0'),
    'unknown': PatternFill('solid', fgColor='E2E8F0'),
}
HEAD_FILL = PatternFill('solid', fgColor='0F4C5C')
THIN = Side(style='thin', color='CBD5E1')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_workbook(data):
    days = data['days']
    wb = Workbook()
    ws = wb.active
    ws.title = data['label'][:31]

    ws['A1'] = f'PEMANTAUAN STOK REAGEN – {data["label"].upper()}'
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = (f'Total {data["total_reagen"]} reagen | Kritis {data["counts"].get("critical", 0)} | '
                f'Waspada {data["counts"].get("warning", 0)} | Aman {data["counts"].get("safe", 0)} | '
                f'Perlu Cek {data["counts"].get("unknown", 0)}')

    head = (['No', 'Nama Reagen', 'Satuan', 'Saldo Awal'] + [str(d) for d in range(1, days + 1)]
            + ['QC', 'Total Pemakaian', 'Stok Masuk', 'Sisa Stok', 'Buffer Stock', 'Status'])
    ws.append([])
    ws.append(head)
    hr = ws.max_row
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=hr, column=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER

    for i, r in enumerate(data['rows'], 1):
        hari = r.get('hari') or {}
        row = ([i, r['nama_reagen'], r.get('satuan') or '', r.get('saldo_awal')]
               + [hari.get(str(d)) or None for d in range(1, days + 1)]
               + [r.get('qc') or 0, r.get('total_pemakaian'), r.get('stok_masuk'), r.get('sisa_stock'),
                  r.get('buffer_stock'), r.get('status_label')])
        ws.append(row)
        rr = ws.max_row
        fill = STATUS_FILL.get(r.get('status'), STATUS_FILL['unknown'])
        for c in range(1, len(head) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.border = BORDER
            if c > 2:
                cell.alignment = Alignment(horizontal='center')
        ws.cell(row=rr, column=2).fill = fill
        ws.cell(row=rr, column=len(head) - 2).fill = fill
        ws.cell(row=rr, column=len(head) - 2).font = Font(bold=True)
        ws.cell(row=rr, column=len(head)).fill = fill
        ws.cell(row=rr, column=len(head)).font = Font(bold=True)

    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 8
    ws.column_dimensions['D'].width = 11
    for d in range(days):
        ws.column_dimensions[get_column_letter(5 + d)].width = 4.5
    for k, w in enumerate([6, 10, 10, 10, 10, 12]):
        ws.column_dimensions[get_column_letter(5 + days + k)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=5)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
